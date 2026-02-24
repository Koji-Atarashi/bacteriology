#!/usr/bin/env python3
import os
import re
import sys
import time
import argparse
import sqlite3
from datetime import datetime

# Optional extractors (installed in your environment in the existing app)
try:
    import nbformat
except Exception:
    nbformat = None

try:
    import pandas as pd
except Exception:
    pd = None

try:
    from docx import Document
except Exception:
    Document = None

try:
    from pptx import Presentation
except Exception:
    Presentation = None

# Optional: minimal web UI
try:
    from flask import Flask, request, render_template_string
except Exception:
    Flask = None


# =========================
# CONFIG
# =========================
ALLOWED_PATHS = ["/mnt/hdd1", "/mnt/e"]  # allowlist
DEFAULT_DB = "/mnt/hdd1/.search_index.sqlite3"

SEARCH_ALLOWED_EXT = {
    ".txt", ".md", ".py", ".json", ".yml", ".yaml", ".sh", ".bash",
    ".r", ".sql", ".html", ".htm", ".css", ".js",
    ".log", ".csv", ".tsv",
    ".docx", ".ipynb", ".xlsx", ".pptx",
}

# Extraction size guards
MAX_BYTES_TEXT = 2 * 1024 * 1024       # 2MB for plain text
MAX_BYTES_DOCX = 20 * 1024 * 1024      # 20MB
MAX_BYTES_IPYNB = 10 * 1024 * 1024     # 10MB
MAX_BYTES_XLSX = 10 * 1024 * 1024      # 10MB
MAX_BYTES_PPTX = 30 * 1024 * 1024      # 30MB

# Indexing guards
DEFAULT_MAX_FILES = 300000
DEFAULT_BATCH = 500

# Exclusions (recommended)
EXCLUDE_DIRS = {
    ".ipynb_checkpoints",
    ".git",
    "__pycache__",
    "node_modules",
    ".venv",
    "venv",
    ".mypy_cache",
    ".pytest_cache",
}

EXCLUDE_FILE_PATTERNS = [
    r"~$",                 # backups
    r"\.tmp$",             # tmp
    r"\.swp$",             # vim
]
_EX_FILE_RE = re.compile("|".join(EXCLUDE_FILE_PATTERNS)) if EXCLUDE_FILE_PATTERNS else None


# =========================
# SAFE PATH
# =========================
def realpath(p: str) -> str:
    return os.path.realpath(p)

def is_safe_path(path: str) -> bool:
    rp = realpath(path)
    for base in ALLOWED_PATHS:
        b = realpath(base)
        if rp == b or rp.startswith(b + os.sep):
            return True
    return False

def normalize_root(root: str) -> str:
    rp = realpath(root)
    if not is_safe_path(rp):
        raise ValueError(f"Root not allowed: {root}")
    if not os.path.isdir(rp):
        raise ValueError(f"Root is not a directory: {root}")
    return rp


# =========================
# FILE TYPE HELPERS
# =========================
def is_probably_binary(path: str, sniff: int = 4096) -> bool:
    try:
        with open(path, "rb") as f:
            chunk = f.read(sniff)
        if b"\x00" in chunk:
            return True
        textchars = bytearray({7, 8, 9, 10, 12, 13, 27} | set(range(0x20, 0x100)))
        nontext = chunk.translate(None, textchars)
        return (len(nontext) / max(1, len(chunk))) > 0.30
    except Exception:
        return True

def get_ext(path: str) -> str:
    return os.path.splitext(path.lower())[1]


# =========================
# TEXT EXTRACTION
# =========================
def read_small_text(path: str, max_bytes: int) -> str | None:
    try:
        if os.path.getsize(path) > max_bytes:
            return None
        if is_probably_binary(path):
            return None
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        return None

def extract_docx(path: str) -> str | None:
    if Document is None:
        return None
    try:
        if os.path.getsize(path) > MAX_BYTES_DOCX:
            return None
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        return None

def extract_ipynb(path: str) -> str | None:
    if nbformat is None:
        return None
    try:
        if os.path.getsize(path) > MAX_BYTES_IPYNB:
            return None
        nb = nbformat.read(path, as_version=4)
        parts = []
        for cell in nb.cells:
            if cell.get("cell_type") in ("markdown", "code"):
                parts.append(cell.get("source", ""))
        return "\n\n".join(parts)
    except Exception:
        return None

def extract_xlsx(path: str) -> str | None:
    if pd is None:
        return None
    try:
        if os.path.getsize(path) > MAX_BYTES_XLSX:
            return None
        xls = pd.ExcelFile(path)
        out = []
        for sheet in xls.sheet_names[:20]:
            df = xls.parse(sheet_name=sheet, dtype=str)
            out.append(f"[SHEET] {sheet}")
            out.append(df.fillna("").astype(str).to_csv(index=False))
        return "\n".join(out)
    except Exception:
        return None

def extract_pptx(path: str) -> str | None:
    if Presentation is None:
        return None
    try:
        if os.path.getsize(path) > MAX_BYTES_PPTX:
            return None
        pres = Presentation(path)
        out = []
        for i, slide in enumerate(pres.slides, start=1):
            out.append(f"[SLIDE {i}]")
            for shp in slide.shapes:
                if hasattr(shp, "text"):
                    t = (shp.text or "").strip()
                    if t:
                        out.append(t)
        return "\n".join(out)
    except Exception:
        return None

def extract_text(path: str) -> str | None:
    ext = get_ext(path)
    if ext not in SEARCH_ALLOWED_EXT:
        return None

    if ext in {
        ".txt", ".md", ".py", ".json", ".yml", ".yaml", ".sh", ".bash",
        ".r", ".sql", ".html", ".htm", ".css", ".js", ".log", ".csv", ".tsv"
    }:
        return read_small_text(path, MAX_BYTES_TEXT)

    if ext == ".docx":
        return extract_docx(path)
    if ext == ".ipynb":
        return extract_ipynb(path)
    if ext == ".xlsx":
        return extract_xlsx(path)
    if ext == ".pptx":
        return extract_pptx(path)

    return None


# =========================
# INDEX (SQLite FTS5)
# =========================
SCHEMA_SQL = """
PRAGMA journal_mode=WAL;
PRAGMA synchronous=NORMAL;

CREATE TABLE IF NOT EXISTS files (
  path TEXT PRIMARY KEY,
  mtime INTEGER NOT NULL,
  size INTEGER NOT NULL
);

CREATE VIRTUAL TABLE IF NOT EXISTS fts USING fts5(path, content, tokenize='unicode61');
"""

def open_db(db_path: str) -> sqlite3.Connection:
    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA foreign_keys=ON;")
    conn.executescript(SCHEMA_SQL)
    return conn

def upsert_document(conn: sqlite3.Connection, path: str, mtime: int, size: int, content: str):
    conn.execute(
        "INSERT INTO files(path, mtime, size) VALUES(?,?,?) "
        "ON CONFLICT(path) DO UPDATE SET mtime=excluded.mtime, size=excluded.size",
        (path, mtime, size),
    )
    # FTS5: safest replacement = DELETE then INSERT
    conn.execute("DELETE FROM fts WHERE path=?", (path,))
    conn.execute("INSERT INTO fts(path, content) VALUES(?, ?)", (path, content))

def remove_document(conn: sqlite3.Connection, path: str):
    conn.execute("DELETE FROM files WHERE path=?", (path,))
    conn.execute("DELETE FROM fts WHERE path=?", (path,))

def iter_files(root: str):
    for base, dirs, files in os.walk(root):
        if not is_safe_path(base):
            continue

        # prune excluded dirs
        dirs[:] = [d for d in dirs if d not in EXCLUDE_DIRS]

        for name in files:
            if _EX_FILE_RE and _EX_FILE_RE.search(name):
                continue
            full = os.path.join(base, name)
            if not is_safe_path(full):
                continue
            yield full

def build_index(db_path: str, roots: list[str], max_files: int, batch: int, delete_missing: bool):
    conn = open_db(db_path)
    cur = conn.cursor()

    seen = set()

    t0 = time.time()
    n_scanned = 0
    n_indexed = 0
    n_skipped = 0

    for root in roots:
        root = normalize_root(root)
        for path in iter_files(root):
            n_scanned += 1
            if n_scanned > max_files:
                print(f"[STOP] reached max_files={max_files}")
                break

            ext = get_ext(path)
            if ext not in SEARCH_ALLOWED_EXT:
                n_skipped += 1
                continue

            try:
                st = os.stat(path)
                mtime = int(st.st_mtime)
                size = int(st.st_size)
            except Exception:
                n_skipped += 1
                continue

            seen.add(path)

            row = cur.execute("SELECT mtime, size FROM files WHERE path=?", (path,)).fetchone()
            if row and row[0] == mtime and row[1] == size:
                continue

            content = extract_text(path)
            if not content:
                if row:
                    remove_document(conn, path)
                n_skipped += 1
                continue

            upsert_document(conn, path, mtime, size, content)
            n_indexed += 1

            if n_indexed % batch == 0:
                conn.commit()
                elapsed = time.time() - t0
                print(f"[COMMIT] indexed={n_indexed} scanned={n_scanned} skipped={n_skipped} elapsed={elapsed:.1f}s")

        if n_scanned > max_files:
            break

    conn.commit()

    if delete_missing:
        all_paths = [r[0] for r in conn.execute("SELECT path FROM files").fetchall()]
        removed = 0
        for p in all_paths:
            if p not in seen:
                remove_document(conn, p)
                removed += 1
        conn.commit()
        print(f"[CLEAN] removed missing={removed}")

    elapsed = time.time() - t0
    print(f"[DONE] db={db_path}")
    print(f"       scanned={n_scanned} indexed={n_indexed} skipped={n_skipped} elapsed={elapsed:.1f}s")


# =========================
# SEARCH (FTS query normalization: 対策2)
# =========================
def normalize_fts_query(user_q: str) -> str:
    """
    対策2:
      - ユーザーが FTS構文(OR/AND/NOT/NEAR/*/"()など)を使っていそうならそのまま尊重
      - そうでなければ "..." のフレーズ検索に変換して安全化（. を含む accession などで落ちない）
    """
    q = (user_q or "").strip()
    if not q:
        return q

    # "FTS構文" っぽいものがある場合はユーザーの意図を尊重
    # 例: foo OR bar, "bile acid", foo*, (foo bar), NEAR
    if any(tok in q for tok in ['"', "*", " NEAR ", " OR ", " AND ", " NOT ", "(", ")"]):
        return q

    # それ以外は安全にフレーズ化
    q = q.replace('"', '""')  # phrase 内の " は "" に
    return f'"{q}"'


def search_index(db_path: str, q: str, limit: int):
    conn = open_db(db_path)
    q2 = normalize_fts_query(q)

    sql = """
    SELECT
      fts.path,
      snippet(fts, 1, '[', ']', '…', 12) AS snip,
      files.mtime,
      files.size
    FROM fts
    JOIN files ON files.path = fts.path
    WHERE fts MATCH ?
    ORDER BY rank
    LIMIT ?;
    """
    return conn.execute(sql, (q2, limit)).fetchall()


# =========================
# Minimal Flask UI
# =========================
HTML = """
<!doctype html>
<html><head><meta charset="utf-8"><title>Search</title>
<style>
body{font-family:-apple-system,system-ui,sans-serif;margin:18px;}
input{padding:8px;width:70%;}
button{padding:8px 10px;}
pre{white-space:pre-wrap;background:#f6f6f6;padding:10px;border-radius:8px;}
.small{color:#666;font-size:12px;}
.err{color:#b00020;background:#fff0f0;padding:10px;border-radius:8px;white-space:pre-wrap;}
</style>
</head><body>
<h2>🔎 Search (Index)</h2>
<form method="get">
  <input name="q" value="{{ q }}" placeholder='FTS query (e.g., microbiota OR bile) or phrase "foo bar"'>
  <button type="submit">Search</button>
</form>
<p class="small">db: {{ db_path }} | now: {{ now }}</p>

{% if err %}
  <div class="err"><b>Error</b>\n{{ err }}</div>
{% endif %}

{% if q and not err %}
  <h3>Hits: {{ results|length }}</h3>
  <ul>
  {% for r in results %}
    <li style="margin:14px 0;">
      <div><b>{{ r.path }}</b></div>
      <div class="small">mtime={{ r.mtime }} size={{ r.size }}</div>
      <pre>{{ r.snip }}</pre>
    </li>
  {% endfor %}
  </ul>
{% endif %}
</body></html>
"""

def serve(db_path: str, host: str, port: int):
    if Flask is None:
        print("Flask is not available in this environment.")
        sys.exit(2)

    app = Flask(__name__)

    @app.route("/", methods=["GET"])
    def home():
        q = (request.args.get("q") or "").strip()
        results = []
        err = None

        if q:
            try:
                rows = search_index(db_path, q, limit=50)
                for path, snip, mtime, size in rows:
                    results.append({
                        "path": path,
                        "snip": snip,
                        "mtime": datetime.fromtimestamp(mtime).isoformat(sep=" ", timespec="seconds"),
                        "size": size
                    })
            except Exception as e:
                err = str(e)

        return render_template_string(
            HTML,
            q=q,
            results=results,
            err=err,
            db_path=db_path,
            now=datetime.now().isoformat(sep=" ", timespec="seconds")
        )

    print(f"Running search UI on http://{host}:{port}  (db={db_path})")
    app.run(host=host, port=port)


# =========================
# CLI
# =========================
def main():
    p = argparse.ArgumentParser(description="Standalone content search + index (SQLite FTS5)")
    sub = p.add_subparsers(dest="cmd", required=True)

    pb = sub.add_parser("build", help="Build/Update index")
    pb.add_argument("--db", default=DEFAULT_DB)
    pb.add_argument("--root", action="append", required=True,
                    help="Root directory to index (repeatable). Must be within ALLOWED_PATHS.")
    pb.add_argument("--max-files", type=int, default=DEFAULT_MAX_FILES)
    pb.add_argument("--batch", type=int, default=DEFAULT_BATCH)
    pb.add_argument("--delete-missing", action="store_true",
                    help="Remove indexed entries not found under provided roots (use carefully).")

    ps = sub.add_parser("search", help="Search the index (CLI)")
    ps.add_argument("--db", default=DEFAULT_DB)
    ps.add_argument("--q", required=True, help="FTS query")
    ps.add_argument("--limit", type=int, default=20)

    pv = sub.add_parser("serve", help="Start minimal web UI for search")
    pv.add_argument("--db", default=DEFAULT_DB)
    pv.add_argument("--host", default="0.0.0.0")
    pv.add_argument("--port", type=int, default=8090)

    args = p.parse_args()

    if args.cmd == "build":
        build_index(args.db, args.root, args.max_files, args.batch, args.delete_missing)

    elif args.cmd == "search":
        rows = search_index(args.db, args.q, args.limit)
        for path, snip, mtime, size in rows:
            ts = datetime.fromtimestamp(mtime).isoformat(sep=" ", timespec="seconds")
            print(f"- {path}  (mtime={ts}, size={size})")
            print(f"  {snip}")
            print()
        print(f"[TOTAL] {len(rows)}")

    elif args.cmd == "serve":
        serve(args.db, args.host, args.port)


if __name__ == "__main__":
    main()
